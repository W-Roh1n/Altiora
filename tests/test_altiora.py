"""Tests for the Altiora interface helpers and Streamlit entry point."""

from __future__ import annotations

import io
import math
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.assumptions import AssumptionError, illustrative_base_case
from src.lbo_model import run_lbo
from src.reporting import (
    assumptions_from_mapping,
    build_workbook,
    case_summary_frame,
    default_sensitivity,
    displayed_outputs,
    evaluate_sensitivity_cell,
    format_irr,
    format_moic,
    format_usd_millions,
    run_standard_cases,
    sources_uses_frame,
    sponsor_cash_flow_frame,
)


def _mapping_from_case(**overrides: object) -> dict[str, object]:
    case = illustrative_base_case()
    payload: dict[str, object] = {
        "company_name": case.company_name,
        "holding_period_years": case.holding_period_years,
        "ltm_revenue": case.ltm_revenue,
        "ltm_ebitda_margin": case.ltm_ebitda_margin,
        "revenue_growth": case.revenue_growth,
        "ebitda_margin": case.ebitda_margin,
        "da_pct_of_revenue": case.da_pct_of_revenue,
        "capex_pct_of_revenue": case.capex_pct_of_revenue,
        "nwc_pct_of_revenue": case.nwc_pct_of_revenue,
        "entry_ebitda_multiple": case.entry_ebitda_multiple,
        "exit_ebitda_multiple": case.exit_ebitda_multiple,
        "senior_leverage": case.senior_leverage,
        "subordinated_leverage": case.subordinated_leverage,
        "senior_interest_rate": case.senior_interest_rate,
        "subordinated_interest_rate": case.subordinated_interest_rate,
        "senior_mandatory_amort_pct": case.senior_mandatory_amort_pct,
        "cash_sweep_pct": case.cash_sweep_pct,
        "tax_rate": case.tax_rate,
        "transaction_fees_pct_of_ev": case.transaction_fees_pct_of_ev,
        "cash_to_balance_sheet": case.cash_to_balance_sheet,
        "minimum_cash": case.minimum_cash,
        "existing_net_debt": case.existing_net_debt,
    }
    payload.update(overrides)
    return payload


def test_mapping_round_trip_matches_engine() -> None:
    original = run_lbo(illustrative_base_case())
    rebuilt = run_lbo(assumptions_from_mapping(_mapping_from_case()))
    assert rebuilt.transaction.entry_enterprise_value == pytest.approx(
        original.transaction.entry_enterprise_value
    )
    assert rebuilt.transaction.sponsor_equity == pytest.approx(original.transaction.sponsor_equity)
    assert rebuilt.exit_enterprise_value == pytest.approx(original.exit_enterprise_value)
    assert rebuilt.gross_moic == pytest.approx(original.gross_moic)
    assert rebuilt.irr == pytest.approx(original.irr)
    assert rebuilt.assumptions.is_illustrative
    assert "not actual" in rebuilt.assumptions.notes.lower()


def test_mapping_labels_unmarked_names_illustrative() -> None:
    case = assumptions_from_mapping(_mapping_from_case(company_name="Northwind Parts"))
    assert case.is_illustrative
    assert "illustrative" in case.company_name.lower()


def test_mapping_rejects_invalid_margin() -> None:
    with pytest.raises(AssumptionError):
        assumptions_from_mapping(_mapping_from_case(ltm_ebitda_margin=0))


def test_changed_entry_multiple_changes_engine_outputs() -> None:
    base = run_lbo(assumptions_from_mapping(_mapping_from_case()))
    higher = run_lbo(assumptions_from_mapping(_mapping_from_case(entry_ebitda_multiple=9.0)))
    assert higher.transaction.entry_enterprise_value == pytest.approx(405.0)
    assert higher.transaction.entry_enterprise_value != pytest.approx(
        base.transaction.entry_enterprise_value
    )
    assert higher.transaction.sponsor_equity != pytest.approx(base.transaction.sponsor_equity)


def test_standard_cases_preserve_scenario_order() -> None:
    cases, errors = run_standard_cases(illustrative_base_case())
    assert errors == {}
    assert list(cases) == ["base", "downside", "upside"]
    assert cases["downside"].irr < cases["base"].irr < cases["upside"].irr


def test_failed_overlay_is_reported_not_filled() -> None:
    base = illustrative_base_case().with_updates(ebitda_margin=(0.012,) * 5)
    cases, errors = run_standard_cases(base)
    assert "base" in cases
    assert "downside" not in cases
    assert "downside" in errors
    assert "EBITDA" in errors["downside"] or "ebitda" in errors["downside"].lower()


def test_sensitivity_cells_match_full_reruns() -> None:
    base = illustrative_base_case()
    irr, moic, notes = default_sensitivity(base)
    assert notes == []
    direct = run_lbo(base)
    assert irr.loc[base.exit_ebitda_multiple, 0.0] == pytest.approx(direct.irr)
    assert moic.loc[base.exit_ebitda_multiple, 0.0] == pytest.approx(direct.gross_moic)

    shocked = base.with_updates(
        exit_ebitda_multiple=9.0,
        revenue_growth=tuple(growth + 0.01 for growth in base.revenue_growth),
    )
    expected = run_lbo(shocked, scenario_name="sensitivity")
    assert irr.loc[9.0, 0.01] == pytest.approx(expected.irr)
    assert moic.loc[9.0, 0.01] == pytest.approx(expected.gross_moic)


def test_sources_uses_and_cash_flows_come_from_result() -> None:
    result = run_lbo(illustrative_base_case())
    frame = sources_uses_frame(result)
    totals = frame.set_index("Item")["Amount"]
    assert totals["Total sources"] == pytest.approx(result.transaction.total_sources)
    assert totals["Total uses"] == pytest.approx(result.transaction.total_uses)
    flows = sponsor_cash_flow_frame(result)
    assert list(flows["Sponsor cash flow"]) == pytest.approx(result.sponsor_cash_flows)
    assert flows.iloc[0]["Sponsor cash flow"] == pytest.approx(-result.transaction.sponsor_equity)
    assert flows.iloc[-1]["Sponsor cash flow"] == pytest.approx(result.exit_equity_value)


def test_workbook_contains_engine_schedules() -> None:
    case = illustrative_base_case()
    result = run_lbo(case)
    payload = build_workbook(case)
    assert payload[:2] == b"PK"
    workbook = load_workbook(io.BytesIO(payload), data_only=True)
    assert {
        "Disclaimer",
        "Assumptions",
        "Sources and Uses",
        "Scenario Comparison",
        "Sponsor Cash Flows",
        "Model Checks",
        "Sensitivity IRR",
        "Operating Base",
        "Debt Base",
        "Cash Base",
        "Operating Downside",
        "Debt Upside",
    } <= set(workbook.sheetnames)
    comparison = workbook["Scenario Comparison"]
    headers = [cell.value for cell in comparison[1]]
    entry_col = headers.index("Entry enterprise value") + 1
    assert comparison.cell(2, entry_col).value == pytest.approx(result.transaction.entry_enterprise_value)
    disclaimer = " ".join(
        str(value) for row in workbook["Disclaimer"].iter_rows(values_only=True) for value in row
    )
    assert "illustrative" in disclaimer.lower()
    assert "investment forecast" in disclaimer.lower()
    summary = case_summary_frame({"base": result})
    assert summary.loc[0, "Gross MOIC"] == pytest.approx(result.gross_moic)


def test_screen_figures_match_workbook_rounding() -> None:
    """The rounded figures on screen are the same engine values stored in Excel."""
    case = illustrative_base_case()
    result = run_lbo(case)
    shown = displayed_outputs(result)
    payload = build_workbook(case)
    workbook = load_workbook(io.BytesIO(payload), data_only=True)
    comparison = workbook["Scenario Comparison"]
    headers = [cell.value for cell in comparison[1]]
    row = {headers[index]: comparison.cell(2, index + 1).value for index in range(len(headers))}
    assert format_usd_millions(row["Entry enterprise value"]) == shown["entry_enterprise_value"]
    assert format_usd_millions(row["Sponsor equity"]) == shown["sponsor_equity"]
    assert format_usd_millions(row["Exit enterprise value"]) == shown["exit_enterprise_value"]
    assert format_usd_millions(row["Exit equity"]) == shown["exit_equity_value"]
    assert format_usd_millions(row["Exit net debt"]) == shown["exit_net_debt"]
    assert format_irr(row["Gross IRR"]) == shown["gross_irr"]
    assert format_moic(row["Gross MOIC"]) == shown["gross_moic"]
    assert shown["entry_enterprise_value"] == "360.00"
    flows = [cell.value for cell in workbook["Sponsor Cash Flows"]["B"]][1:]
    assert flows == pytest.approx(result.sponsor_cash_flows)


def test_rejected_sensitivity_cell_is_not_invented() -> None:
    irr, moic, note = evaluate_sensitivity_cell(illustrative_base_case(), -1.0, 0.0)
    assert math.isnan(irr)
    assert math.isnan(moic)
    assert note is not None
    assert "positive" in note.lower()


def test_entry_value_must_exceed_existing_net_debt() -> None:
    case = illustrative_base_case().with_updates(existing_net_debt=10_000.0)
    with pytest.raises(ValueError, match="existing net debt"):
        run_lbo(case)


def test_zero_senior_leverage_still_balances() -> None:
    base = illustrative_base_case()
    unlevered = run_lbo(base.with_updates(senior_leverage=0.0, senior_mandatory_amort_pct=0.0))
    levered = run_lbo(base)
    assert unlevered.transaction.new_senior_debt == pytest.approx(0.0)
    assert unlevered.transaction.total_sources == pytest.approx(unlevered.transaction.total_uses)
    assert unlevered.transaction.sponsor_equity > levered.transaction.sponsor_equity
    assert float(unlevered.debt_schedule.loc["Closing senior", "Year 5"]) == pytest.approx(0.0)


def test_monster_filing_anchors_are_not_treated_as_unlevered() -> None:
    """Historical debt of zero still leaves room for hypothetical acquisition debt."""
    from src.filings import (
        FILING_URL,
        SHARES_OUTSTANDING,
        derived_ebitda,
        engine_ebitda_margin,
        filing_register,
        income_statement_ties,
        monster_hypothetical_case,
        nonaffiliate_equity_usd_millions,
        reported_cost_of_sales,
        reported_net_income,
        reported_net_sales,
        reported_operating_expenses,
        reported_operating_income,
        reported_depreciation,
    )

    case = monster_hypothetical_case()
    assert case.is_illustrative
    assert case.ltm_revenue == pytest.approx(reported_net_sales(2025))
    assert reported_net_sales(2025) == pytest.approx(8_294.343)
    assert reported_net_sales(2024) == pytest.approx(7_492.709)
    assert reported_net_sales(2023) == pytest.approx(7_140.027)
    assert reported_cost_of_sales(2025) == pytest.approx(3_662.148)
    assert reported_operating_expenses(2025) == pytest.approx(2_212.841)
    assert reported_net_income(2025) == pytest.approx(1_905.432)
    assert all(income_statement_ties(year) for year in (2025, 2024, 2023))
    assert case.entry_ebitda_multiple == pytest.approx(15.0)
    shares = next(row for row in filing_register() if row.item == "Common shares outstanding")
    assert shares.amount_usd_millions is None
    assert f"{SHARES_OUTSTANDING:,}" in shares.note
    market = next(
        row
        for row in filing_register()
        if row.item.startswith("Aggregate market value")
    )
    assert market.amount_usd_millions == pytest.approx(nonaffiliate_equity_usd_millions())
    assert market.amount_usd_millions == pytest.approx(56_505.47148)
    assert "not enterprise value" in market.note.lower() or "not used as the entry multiple" in market.note.lower()
    assert derived_ebitda(2025) == pytest.approx(
        reported_operating_income(2025) + reported_depreciation(2025)
    )
    assert derived_ebitda(2025) == pytest.approx(2_533.795)
    assert case.ltm_ebitda_margin == pytest.approx(engine_ebitda_margin(2025))
    assert case.ltm_ebitda == pytest.approx(case.ltm_revenue * case.ltm_ebitda_margin)
    assert abs(case.ltm_ebitda - derived_ebitda(2025)) < 0.01
    assert case.existing_net_debt == pytest.approx(0.0)
    assert case.senior_leverage > 0
    assert case.subordinated_leverage > 0

    result = run_lbo(case)
    assert result.transaction.refinanced_net_debt == pytest.approx(0.0)
    assert result.transaction.new_senior_debt == pytest.approx(case.senior_leverage * case.ltm_ebitda)
    assert result.transaction.new_subordinated_debt > 0
    assert result.transaction.equity_purchase_price == pytest.approx(result.transaction.entry_enterprise_value)
    assert result.transaction.total_sources == pytest.approx(result.transaction.total_uses)
    assert result.transaction.sponsor_equity > 0
    assert result.exit_equity_value == pytest.approx(
        max(0.0, result.exit_enterprise_value - result.exit_net_debt)
    )
    assert math.isfinite(result.irr)
    assert result.gross_moic > 0

    blank = [row for row in filing_register() if row.status == "unverified"]
    assert blank
    assert all(row.amount_usd_millions is None for row in blank)
    assert all(row.note.strip() for row in blank)

    payload = build_workbook(case)
    workbook = load_workbook(io.BytesIO(payload), data_only=True)
    register = " ".join(
        str(value)
        for row in workbook["Filing register"].iter_rows(values_only=True)
        for value in row
        if value is not None
    )
    assert FILING_URL in register
    assert "8,294.343" in register or "8294.343" in register
    assert "unverified" in register
    disclaimer = " ".join(
        str(value) for row in workbook["Disclaimer"].iter_rows(values_only=True) for value in row
    )
    assert "investment forecast" in disclaimer.lower()
    assert "fictional" not in disclaimer.lower()
    shown = displayed_outputs(result)
    comparison = workbook["Scenario Comparison"]
    headers = [cell.value for cell in comparison[1]]
    row = {headers[index]: comparison.cell(2, index + 1).value for index in range(len(headers))}
    assert format_usd_millions(row["Entry enterprise value"]) == shown["entry_enterprise_value"]
    assert format_usd_millions(row["Sponsor equity"]) == shown["sponsor_equity"]
    assert format_irr(row["Gross IRR"]) == shown["gross_irr"]
    assert format_moic(row["Gross MOIC"]) == shown["gross_moic"]

    higher = run_lbo(case.with_updates(senior_leverage=case.senior_leverage + 0.5))
    assert higher.transaction.new_senior_debt > result.transaction.new_senior_debt
    assert higher.transaction.sponsor_equity < result.transaction.sponsor_equity
    assert higher.transaction.entry_enterprise_value == pytest.approx(
        result.transaction.entry_enterprise_value
    )


def test_app_does_not_reimplement_financial_formulas() -> None:
    source = (ROOT / "app.py").read_text(encoding="utf-8")
    assert "Altiora" in source
    assert "run_lbo" in source
    assert "build_workbook" in source
    for forbidden in ("ltm_ebitda *", "annualized_irr(", "gross_moic(", "exit_enterprise_value("):
        assert forbidden not in source


def test_streamlit_app_recalculates_entry_value() -> None:
    from streamlit.testing.v1 import AppTest

    from src.filings import monster_hypothetical_case

    case = monster_hypothetical_case()
    shown = displayed_outputs(run_lbo(case))
    higher = displayed_outputs(
        run_lbo(case.with_updates(entry_ebitda_multiple=case.entry_ebitda_multiple + 1.0))
    )
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
    app.run()
    assert not app.exception
    overview = _visible_text(app)
    assert "Altiora" in overview
    assert "Monster Beverage" in overview
    assert shown["entry_enterprise_value"] in overview
    assert "Exit equity" in overview
    assert "illustrative" in overview.lower()
    assert "investment advice" in overview.lower()
    assert "Unverified" in overview
    assert "sec.gov" in overview

    app.radio(key="nav_page").set_value("Deal Assumptions").run()
    assert not app.exception
    app.number_input(key="ctrl_entry_multiple").set_value(case.entry_ebitda_multiple + 1.0).run()
    assert not app.exception
    recalculated = _visible_text(app)
    assert higher["entry_enterprise_value"] in recalculated
    assert shown["entry_enterprise_value"] not in recalculated


def test_streamlit_invalid_cash_warns_and_keeps_last_case() -> None:
    from streamlit.testing.v1 import AppTest

    from src.filings import monster_hypothetical_case

    shown = displayed_outputs(run_lbo(monster_hypothetical_case()))
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
    app.run()
    app.radio(key="nav_page").set_value("Deal Assumptions").run()
    app.number_input(key="ctrl_min_cash").set_value(5_000.0).run()
    assert not app.exception
    page = _visible_text(app)
    assert "minimum cash" in page.lower()
    assert shown["entry_enterprise_value"] in page


def _visible_text(app: object) -> str:
    chunks: list[str] = []
    for name in ("markdown", "header", "subheader", "caption", "metric", "success", "error", "warning", "text"):
        for element in getattr(app, name):
            chunks.append(str(getattr(element, "value", "") or ""))
            chunks.append(str(getattr(element, "label", "") or ""))
    return "\n".join(chunks)
