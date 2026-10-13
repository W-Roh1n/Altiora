"""Presentation helpers that call the LBO engine.

No financial formulas live here. Schedules, returns, sensitivity cells,
and workbook values are outputs of ``run_lbo``.
"""

from __future__ import annotations

import io
import math
from typing import Mapping

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.assumptions import (
    STANDARD_SCENARIOS,
    AssumptionError,
    DealAssumptions,
    apply_scenario,
)
from src.filings import FILING_URL, filing_anchor_gaps, register_frame_records
from src.lbo_model import LBOResult, run_lbo
from src.returns import ReturnsError
from src.validation import ValidationError

MODEL_ERRORS = (AssumptionError, ValueError, ValidationError, ReturnsError, RuntimeError)

ILLUSTRATIVE_NOTE = (
    "ILLUSTRATIVE HYPOTHETICAL ONLY — not actual company financials. "
    "Figures are fictional and are not an investment forecast."
)

_INPUT_UNITS = {
    "company_name": "text",
    "is_illustrative": "flag",
    "currency": "text",
    "units": "text",
    "ltm_revenue": "USD millions",
    "ltm_ebitda_margin": "decimal",
    "revenue_growth": "decimal per year",
    "ebitda_margin": "decimal per year",
    "da_pct_of_revenue": "decimal",
    "capex_pct_of_revenue": "decimal",
    "nwc_pct_of_revenue": "decimal",
    "entry_ebitda_multiple": "multiple",
    "exit_ebitda_multiple": "multiple",
    "holding_period_years": "years",
    "senior_leverage": "multiple of LTM EBITDA",
    "subordinated_leverage": "multiple of LTM EBITDA",
    "senior_interest_rate": "decimal",
    "subordinated_interest_rate": "decimal",
    "senior_mandatory_amort_pct": "decimal",
    "cash_sweep_pct": "decimal",
    "tax_rate": "decimal",
    "transaction_fees_pct_of_ev": "decimal",
    "cash_to_balance_sheet": "USD millions",
    "minimum_cash": "USD millions",
    "existing_net_debt": "USD millions",
    "notes": "text",
    "ltm_ebitda": "USD millions",
}

_SCENARIO_ORDER = ("base", "downside", "upside")
_PERCENT_SHEETS = {"Sensitivity IRR"}
_MULTIPLE_SHEETS = {"Sensitivity MOIC"}


def _notes_for(company_name: str) -> str:
    if "monster beverage" in company_name.lower():
        return (
            "HYPOTHETICAL LBO labeled illustrative, not actual reported results and "
            "not an investment forecast. FY2025 filing anchors are listed on the "
            "Filing register. Acquisition debt is an assumption and is not historical debt."
        )
    return ILLUSTRATIVE_NOTE


def assumptions_from_mapping(values: Mapping[str, object]) -> DealAssumptions:
    """Build validated assumptions from control values.

    Rates and margins must already be decimals. Currency stays USD millions:
    the engine does not convert currencies. The case is always marked
    illustrative.
    """
    growth = tuple(float(x) for x in values["revenue_growth"])  # type: ignore[index]
    margins = tuple(float(x) for x in values["ebitda_margin"])  # type: ignore[index]
    holding = int(values["holding_period_years"])
    name = str(values["company_name"]).strip()
    if "illustrative" not in name.lower():
        name = f"{name} (illustrative)"
    return DealAssumptions(
        company_name=name,
        is_illustrative=True,
        currency="USD",
        units="millions",
        ltm_revenue=float(values["ltm_revenue"]),
        ltm_ebitda_margin=float(values["ltm_ebitda_margin"]),
        revenue_growth=growth,
        ebitda_margin=margins,
        da_pct_of_revenue=float(values["da_pct_of_revenue"]),
        capex_pct_of_revenue=float(values["capex_pct_of_revenue"]),
        nwc_pct_of_revenue=float(values["nwc_pct_of_revenue"]),
        entry_ebitda_multiple=float(values["entry_ebitda_multiple"]),
        exit_ebitda_multiple=float(values["exit_ebitda_multiple"]),
        holding_period_years=holding,
        senior_leverage=float(values["senior_leverage"]),
        subordinated_leverage=float(values["subordinated_leverage"]),
        senior_interest_rate=float(values["senior_interest_rate"]),
        subordinated_interest_rate=float(values["subordinated_interest_rate"]),
        senior_mandatory_amort_pct=float(values["senior_mandatory_amort_pct"]),
        cash_sweep_pct=float(values["cash_sweep_pct"]),
        tax_rate=float(values["tax_rate"]),
        transaction_fees_pct_of_ev=float(values["transaction_fees_pct_of_ev"]),
        cash_to_balance_sheet=float(values["cash_to_balance_sheet"]),
        minimum_cash=float(values["minimum_cash"]),
        existing_net_debt=float(values["existing_net_debt"]),
        notes=_notes_for(name),
    )


def run_standard_cases(
    base: DealAssumptions,
) -> tuple[dict[str, LBOResult], dict[str, str]]:
    """Run base, downside, and upside. Failed overlays are reported, not filled in."""
    results: dict[str, LBOResult] = {}
    errors: dict[str, str] = {}
    for key in _SCENARIO_ORDER:
        adj = STANDARD_SCENARIOS[key]
        try:
            case = base if key == "base" else apply_scenario(base, adj)
            results[key] = run_lbo(case, scenario_name=key)
        except MODEL_ERRORS as exc:
            errors[key] = f"{type(exc).__name__}: {exc}"
    return results, errors


def format_usd_millions(value: float, decimals: int = 2) -> str:
    """Display an engine amount. Does not change the stored result."""
    number = float(value)
    if math.isnan(number):
        return "—"
    text = f"{abs(number):,.{decimals}f}"
    if number < 0:
        return f"({text})"
    return text


def format_irr(value: float) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return "n/m"
    return f"{float(value):.1%}"


def format_moic(value: float) -> str:
    return f"{float(value):.2f}x"


def displayed_outputs(result: LBOResult) -> dict[str, str]:
    """The figures the app shows for one calculated case, taken from ``run_lbo``."""
    return {
        "entry_enterprise_value": format_usd_millions(result.transaction.entry_enterprise_value),
        "sponsor_equity": format_usd_millions(result.transaction.sponsor_equity),
        "exit_enterprise_value": format_usd_millions(result.exit_enterprise_value),
        "exit_equity_value": format_usd_millions(result.exit_equity_value),
        "exit_net_debt": format_usd_millions(result.exit_net_debt),
        "gross_irr": format_irr(result.irr),
        "gross_moic": format_moic(result.gross_moic),
    }


def evaluate_sensitivity_cell(
    base: DealAssumptions,
    exit_multiple: float,
    growth_delta: float,
) -> tuple[float, float, str | None]:
    """One full model run. A rejected case returns blank IRR and MOIC, not an estimate."""
    try:
        case = base.with_updates(
            revenue_growth=tuple(growth + growth_delta for growth in base.revenue_growth),
            exit_ebitda_multiple=exit_multiple,
        )
        result = run_lbo(case, scenario_name="sensitivity")
    except MODEL_ERRORS as exc:
        return float("nan"), float("nan"), f"{type(exc).__name__}: {exc}"
    return result.irr, result.gross_moic, None


def default_sensitivity(
    base: DealAssumptions,
) -> tuple[pd.DataFrame, pd.DataFrame, list[str]]:
    """Recalculate IRR and MOIC across exit multiple and revenue-growth shocks.

    Each cell is a full ``run_lbo`` run. Cells the engine rejects are blank
    and listed in the notes. Nothing is interpolated.
    """
    exit_multiples: list[float] = []
    for step in (-2.0, -1.0, 0.0, 1.0, 2.0):
        multiple = round(float(base.exit_ebitda_multiple) + step, 4)
        if multiple > 0 and multiple not in exit_multiples:
            exit_multiples.append(multiple)

    growth_deltas: list[float] = []
    for delta in (-0.02, -0.01, 0.0, 0.01, 0.02):
        shocked = tuple(g + delta for g in base.revenue_growth)
        if all(-0.5 < g < 1.0 for g in shocked):
            growth_deltas.append(delta)

    irr_rows: list[dict[str, float]] = []
    moic_rows: list[dict[str, float]] = []
    notes: list[str] = []
    for multiple in exit_multiples:
        irr_row: dict[str, float] = {"Exit multiple": multiple}
        moic_row: dict[str, float] = {"Exit multiple": multiple}
        for delta in growth_deltas:
            irr_value, moic_value, error = evaluate_sensitivity_cell(base, multiple, delta)
            irr_row[delta] = irr_value
            moic_row[delta] = moic_value
            if error:
                notes.append(f"Exit {multiple:.2f}x, growth {delta * 100:+.1f} pp: {error}")
        irr_rows.append(irr_row)
        moic_rows.append(moic_row)

    irr = pd.DataFrame(irr_rows).set_index("Exit multiple")
    moic = pd.DataFrame(moic_rows).set_index("Exit multiple")
    irr.index.name = "Exit multiple"
    moic.index.name = "Exit multiple"
    return irr, moic, notes


def assumptions_table(assumptions: DealAssumptions) -> pd.DataFrame:
    """List engine inputs, plus LTM EBITDA from ``DealAssumptions.ltm_ebitda``."""
    rows: list[dict[str, object]] = []
    for field_name, unit in _INPUT_UNITS.items():
        if field_name == "ltm_ebitda":
            value: object = assumptions.ltm_ebitda
        else:
            value = getattr(assumptions, field_name)
        if isinstance(value, tuple):
            value = ", ".join(f"{item:.6f}" for item in value)
        rows.append({"Field": field_name, "Unit": unit, "Value": value})
    return pd.DataFrame(rows)


def sources_uses_frame(result: LBOResult) -> pd.DataFrame:
    sources = result.sources.rename("Amount").rename_axis("Item").reset_index()
    sources.insert(0, "Side", "Sources")
    uses = result.uses.rename("Amount").rename_axis("Item").reset_index()
    uses.insert(0, "Side", "Uses")
    totals = pd.DataFrame(
        [
            {
                "Side": "Sources",
                "Item": "Total sources",
                "Amount": result.transaction.total_sources,
            },
            {
                "Side": "Uses",
                "Item": "Total uses",
                "Amount": result.transaction.total_uses,
            },
        ]
    )
    return pd.concat([sources, uses, totals], ignore_index=True)


def sponsor_cash_flow_frame(result: LBOResult) -> pd.DataFrame:
    periods = ["Entry (t=0)", *result.assumptions.year_labels]
    flows = list(result.sponsor_cash_flows)
    if len(periods) != len(flows):
        raise ValueError("Sponsor cash-flow length does not match the holding period")
    return pd.DataFrame({"Period": periods, "Sponsor cash flow": flows})


def case_summary_frame(results: Mapping[str, LBOResult]) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for key in _SCENARIO_ORDER:
        result = results.get(key)
        if result is None:
            continue
        exit_year = result.assumptions.year_labels[-1]
        rows.append(
            {
                "Scenario": key,
                "Entry enterprise value": result.transaction.entry_enterprise_value,
                "Sponsor equity": result.transaction.sponsor_equity,
                "Exit-year revenue": float(result.operating_schedule.loc["Revenue", exit_year]),
                "Exit-year EBITDA": float(result.operating_schedule.loc["EBITDA", exit_year]),
                "Exit enterprise value": result.exit_enterprise_value,
                "Exit equity": result.exit_equity_value,
                "Exit net debt": result.exit_net_debt,
                "Gross MOIC": result.gross_moic,
                "Gross IRR": result.irr,
                "Amortization shortfall": result.financing.total_amortization_shortfall,
                "Min-cash breach": result.financing.total_min_cash_breach,
                "Financing shortfall": result.financing.has_shortfall,
            }
        )
    return pd.DataFrame(rows)


def checks_frame(result: LBOResult) -> pd.DataFrame:
    rows = [
        {
            "Check": check.name,
            "Status": "Pass" if check.passed else "Flag",
            "Detail": check.detail,
        }
        for check in result.checks
    ]
    for note in result.financing.notes:
        rows.append({"Check": "financing_note", "Status": "Flag", "Detail": note})
    return pd.DataFrame(rows)


def _schedule_frame(schedule: pd.DataFrame) -> pd.DataFrame:
    frame = schedule.copy()
    frame.index.name = "Line item"
    return frame.reset_index()


def _labeled_sensitivity(frame: pd.DataFrame, kind: str) -> pd.DataFrame:
    view = frame.copy()
    view.index = [f"{float(idx):.1f}x" for idx in view.index]
    view.columns = [f"{float(col) * 100:+.1f} pp" for col in view.columns]
    view.index.name = "Exit EV / EBITDA"
    view.columns.name = f"{kind} by revenue-growth shock"
    return view.reset_index()


def _disclaimer_text(base: DealAssumptions) -> str:
    """Keep fictional cases and the Monster filing case from using the same claim."""
    if "monster beverage" in base.company_name.lower():
        gaps = filing_anchor_gaps(base)
        anchor = (
            "Some operating inputs no longer match the stored FY2025 anchors."
            if gaps
            else "LTM net sales, the derived EBITDA margin, D&A, capex, trade NWC, "
            "the 23.2% tax rate, and zero existing net debt match the stored FY2025 anchors."
        )
        return (
            "Hypothetical LBO, labeled illustrative, and not an investment forecast. "
            f"{anchor} Historical lines and the source URL are on the Filing register sheet: "
            f"{FILING_URL} Acquisition debt, multiples, growth, fees, and closing cash are "
            "assumptions. They are not Monster's historical borrowings or reported results."
        )
    return (
        "Illustrative hypotheticals only. The company is fictional. "
        "These figures are not historical results, a valuation of a "
        "real business, or an investment forecast. "
        "The Filing register sheet is Monster Beverage's FY2025 Form 10-K and is not "
        "the source of this case."
    )


def build_workbook(base: DealAssumptions) -> bytes:
    """Excel workbook of assumptions and engine schedules for this base case."""
    results, errors = run_standard_cases(base)
    if "base" not in results:
        detail = errors.get("base", "base case failed")
        raise ValueError(detail)

    base_result = results["base"]
    disclaimer = pd.DataFrame(
        [
            {"Item": "Product", "Detail": "Altiora — personal LBO model"},
            {
                "Item": "Notice",
                "Detail": _disclaimer_text(base),
            },
            {"Item": "Units", "Detail": "USD millions. Rates in the assumptions sheet are decimals."},
            {"Item": "Company", "Detail": base.company_name},
            {"Item": "Base scenario", "Detail": base_result.scenario_name},
            {
                "Item": "Scenario errors",
                "Detail": "; ".join(f"{key}: {msg}" for key, msg in errors.items()) or "None",
            },
        ]
    )
    sensitivity_irr, sensitivity_moic, sensitivity_notes = default_sensitivity(base)
    note_frame = pd.DataFrame(
        {"Notice": sensitivity_notes or ["Every sensitivity cell was calculated by the LBO engine."]}
    )
    error_frame = pd.DataFrame(
        [{"Scenario": key, "Error": msg} for key, msg in errors.items()]
        or [{"Scenario": "", "Error": "No scenario failed validation."}]
    )

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        disclaimer.to_excel(writer, sheet_name="Disclaimer", index=False)
        pd.DataFrame(register_frame_records()).to_excel(writer, sheet_name="Filing register", index=False)
        assumptions_table(base).to_excel(writer, sheet_name="Assumptions", index=False)
        sources_uses_frame(base_result).to_excel(writer, sheet_name="Sources and Uses", index=False)
        case_summary_frame(results).to_excel(writer, sheet_name="Scenario Comparison", index=False)
        error_frame.to_excel(writer, sheet_name="Scenario Errors", index=False)
        sponsor_cash_flow_frame(base_result).to_excel(writer, sheet_name="Sponsor Cash Flows", index=False)
        checks_frame(base_result).to_excel(writer, sheet_name="Model Checks", index=False)
        _labeled_sensitivity(sensitivity_irr, "Gross IRR").to_excel(
            writer, sheet_name="Sensitivity IRR", index=False
        )
        _labeled_sensitivity(sensitivity_moic, "Gross MOIC").to_excel(
            writer, sheet_name="Sensitivity MOIC", index=False
        )
        note_frame.to_excel(writer, sheet_name="Sensitivity Notes", index=False)
        for key, result in results.items():
            title = key.title()
            _schedule_frame(result.operating_schedule).to_excel(
                writer, sheet_name=f"Operating {title}", index=False
            )
            _schedule_frame(result.debt_schedule).to_excel(
                writer, sheet_name=f"Debt {title}", index=False
            )
            _schedule_frame(result.cash_schedule).to_excel(
                writer, sheet_name=f"Cash {title}", index=False
            )
        _style_workbook(writer.book)
    return buffer.getvalue()


def _style_workbook(workbook) -> None:
    header_fill = PatternFill("solid", fgColor="0E1C2F")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    body_font = Font(name="Calibri", size=10, color="142033")
    line = Side(style="thin", color="E3E8EE")
    grid = Border(left=line, right=line, top=line, bottom=line)
    for sheet in workbook.worksheets:
        _style_sheet(sheet, header_fill, header_font, body_font, grid)


def _style_sheet(sheet: Worksheet, header_fill, header_font, body_font, grid: Border) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.paperSize = sheet.PAPERSIZE_TABLOID
    sheet.sheet_properties.tabColor = "1F4E79"
    sheet.row_dimensions[1].height = 22
    sheet.sheet_view.showGridLines = False
    percent_sheet = sheet.title in _PERCENT_SHEETS
    multiple_sheet = sheet.title in _MULTIPLE_SHEETS

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=sheet.max_column):
        for cell in row:
            cell.font = body_font
            cell.border = grid
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            header = str(sheet.cell(1, cell.column).value or "")
            if isinstance(cell.value, float):
                if percent_sheet or "IRR" in header:
                    cell.number_format = "0.0%"
                elif multiple_sheet or "MOIC" in header:
                    cell.number_format = '0.00"x"'
                else:
                    cell.number_format = "#,##0.00"

    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        width = 14
        for cell in column:
            if cell.value is not None:
                width = max(width, min(52, len(str(cell.value)) + 3))
        sheet.column_dimensions[letter].width = width

    if sheet.title == "Disclaimer":
        sheet.column_dimensions["B"].width = 92
        for index in range(2, sheet.max_row + 1):
            sheet.row_dimensions[index].height = 36
